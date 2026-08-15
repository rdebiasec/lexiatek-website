#!/usr/bin/env python3
"""Borrador de propuesta de asociación LexiaTek (docx). Salida local, no va a Git."""

from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = Path(__file__).resolve().parent / "LexiaTek_Propuesta_de_Asociacion.docx"

NAVY = RGBColor(0x0A, 0x1F, 0x3D)
TEAL = RGBColor(0x0E, 0x7C, 0x7B)
GOLD = RGBColor(0xC4, 0xA3, 0x5A)
INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x5B, 0x6B, 0x7C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

FECHA = date(2026, 8, 14).strftime("%d de agosto de %Y")


def set_run_font(run, name="Calibri", size=11, bold=False, italic=False, color=INK):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color


def shade_cell(cell, hex_color: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_border(cell, color="D5DEE8"):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), color)
        tcBorders.append(el)
    tcPr.append(tcBorders)


def para(doc, text, *, size=11, bold=False, italic=False, color=INK,
         align="left", space_after=10, space_before=0):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    p.alignment = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }[align]
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic, color=color)
    return p


def h1(doc, text):
    p = para(doc, text, size=16, bold=True, color=NAVY, space_before=18, space_after=8)
    p.paragraph_format.keep_with_next = True
    return p


def h2(doc, text):
    p = para(doc, text, size=13, bold=True, color=TEAL, space_before=14, space_after=6)
    p.paragraph_format.keep_with_next = True
    return p


def body(doc, text):
    return para(doc, text, size=11, align="justify", space_after=10)


def note(doc, text):
    return para(doc, text, size=10, italic=True, color=TEAL, space_after=10)


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.clear()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.left_indent = Cm(1.0)
    run = p.add_run(text)
    set_run_font(run, size=11, color=INK)
    return p


def add_hr(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(12)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "C4A35A")
    pBdr.append(bottom)
    pPr.append(pBdr)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.autofit = True
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        run = p.add_run(h)
        set_run_font(run, size=10, bold=True, color=WHITE)
        shade_cell(cell, "0A1F3D")
        set_cell_border(cell, "0A1F3D")
    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row):
            cell = table.rows[r_i + 1].cells[c_i]
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(val)
            set_run_font(run, size=10, color=INK)
            if r_i % 2 == 1:
                shade_cell(cell, "F4F7FA")
            set_cell_border(cell)
    doc.add_paragraph().paragraph_format.space_after = Pt(8)
    return table


def set_narrow_margins(section):
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)


def add_page_number(section):
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("LexiaTek  ·  BORRADOR para revisión  ·  Confidencial  ·  ")
    set_run_font(run, size=8, color=MUTED)
    fld1 = OxmlElement("w:fldChar")
    fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    r2 = p.add_run()
    r2._r.append(fld1)
    r2._r.append(instr)
    r2._r.append(fld2)
    set_run_font(r2, size=8, color=MUTED)


def build():
    doc = Document()
    section = doc.sections[0]
    set_narrow_margins(section)
    add_page_number(section)
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.font.color.rgb = INK

    para(doc, "BORRADOR PARA REVISIÓN  ·  CONFIDENCIAL", size=10, bold=True, color=GOLD, align="center", space_after=4)
    para(doc, "LEXIATEK", size=28, bold=True, color=NAVY, align="center", space_before=36, space_after=4)
    para(doc, "Propuesta de asociación", size=18, color=TEAL, align="center", space_after=6)
    para(
        doc,
        "Documento de trabajo. Michele Aguilar puede corregir el lenguaje, tachar y anotar.",
        size=11, italic=True, color=MUTED, align="center", space_after=14,
    )
    add_hr(doc)
    para(
        doc,
        "Constitución y operación de un despacho de orientación y representación jurídica en Colombia, "
        "con apoyo tecnológico y con revisión y firma humanas.",
        size=12, italic=True, color=MUTED, align="center", space_after=28,
    )
    para(doc, "Presentada por", size=10, color=MUTED, align="center", space_after=2)
    para(doc, "Ricardo De Biase", size=14, bold=True, color=NAVY, align="center", space_after=0)
    para(doc, "Ingeniero  ·  socio único de DBX Solutions", size=11, color=TEAL, align="center", space_after=18)
    para(doc, "Dirigida a", size=10, color=MUTED, align="center", space_after=2)
    para(doc, "Michele Aguilar", size=14, bold=True, color=NAVY, align="center", space_after=2)
    para(
        doc,
        "en su calidad de futura socia y abogada titular de LexiaTek",
        size=11, color=MUTED, align="center", space_after=36,
    )
    para(doc, "[Ciudad], " + FECHA, size=11, color=INK, align="center", space_after=6)
    para(
        doc,
        "Borrador. No constituye contrato, estatutos ni dictamen jurídico.\n"
        "Ábralo en Microsoft Word o Google Docs, active Control de cambios y reescriba lo que no le represente.",
        size=9, italic=True, color=MUTED, align="center",
    )

    doc.add_page_break()

    para(doc, "[Ciudad], " + FECHA, size=11, space_after=16)
    para(doc, "Estimada doctora Michele Aguilar:", size=11, space_after=8)
    note(
        doc,
        "Cómo usar este borrador: en Word, ficha Revisar → Control de cambios. "
        "Los textos entre [corchetes] son campos pendientes (cédula, tarjeta profesional, domicilio, NIT). "
        "Si un párrafo no le representa, sustitúyalo.",
    )
    body(
        doc,
        "Me dirijo a usted para presentarle, de manera formal, la propuesta de asociación entre Ricardo De Biase "
        "y Michele Aguilar, destinada a constituir y operar LexiaTek como despacho de orientación y representación "
        "jurídica en Colombia.",
    )
    body(
        doc,
        "LexiaTek será de ambos, en partes iguales. Usted será la abogada titular del despacho. "
        "Yo seré el socio técnico. No ejerzo la profesión de abogado, no firmo, no litigo y no me anuncio como tal.",
    )
    body(
        doc,
        "La propiedad intelectual y la arquitectura de la inteligencia artificial pertenecen de manera exclusiva "
        "a DBX Solutions, sociedad de mi única titularidad. Esa tecnología se licencia a usted y, a través de usted, "
        "a LexiaTek, únicamente mientras subsista la asociación. De manera recíproca, sus servicios profesionales "
        "se autorizan a LexiaTek únicamente mientras subsista la misma asociación. Al terminar la asociación, "
        "ambas autorizaciones cesan, sin perjuicio del periodo transitorio que usted confirme para no abandonar clientes.",
    )
    body(
        doc,
        "Este escrito no es un contrato. Es la base para su examen. Usted podrá anotarlo, objetarlo o complementarlo. "
        "Si lo aprueba en lo sustancial, se redactarán los instrumentos definitivos bajo su revisión.",
    )
    body(
        doc,
        "El catálogo de tarifas al público no forma parte de esta propuesta. Se elaborará después, de común acuerdo, "
        "y usted conservará la última palabra sobre cualquier cobro que no resulte éticamente defendible.",
    )
    para(doc, "Quedo a su disposición.", size=11, space_before=8, space_after=16)
    para(doc, "Atentamente,", size=11, space_after=28)
    para(doc, "Ricardo De Biase", size=12, bold=True, color=NAVY, space_after=0)
    para(doc, "Ingeniero", size=11, color=MUTED, space_after=0)
    para(doc, "DBX Solutions", size=11, color=MUTED, space_after=0)

    doc.add_page_break()

    h1(doc, "1. Objeto")
    body(
        doc,
        "La presente propuesta tiene por objeto formalizar la asociación entre Ricardo De Biase y Michele Aguilar "
        "para constituir y operar LexiaTek, un despacho dedicado a la orientación y la representación jurídica "
        "de personas en Colombia que encuentran barreras de costo, de información o de acceso frente a trámites "
        "y procesos frecuentes.",
    )
    body(
        doc,
        "El despacho no se limita al derecho penal. Atenderá, entre otros, asuntos penales; civiles; laborales; "
        "de acoso sexual; peticiones, quejas, reclamos y sugerencias (PQRS); acciones de tutela; y demás procesos "
        "y trámites habituales de la ciudadanía —derechos de petición, impulsos, reclamaciones y gestiones análogas—, "
        "en el alcance que usted, como abogada titular, acepte en cada caso.",
    )
    body(
        doc,
        "LexiaTek se apoyará en herramientas de inteligencia artificial de propiedad exclusiva de DBX Solutions. "
        "Dichas herramientas preparan insumos internos —cronologías, hipótesis preliminares, borradores y control "
        "de términos— para que usted los revise, los ajuste y, si corresponde, los apruebe o los firme. "
        "Ninguna herramienta sustituye el criterio profesional, radica escritos ni se dirige al cliente en calidad de abogado.",
    )

    h1(doc, "2. Las partes")
    h2(doc, "2.1. En LexiaTek")
    body(
        doc,
        "LexiaTek tendrá dos socios, en partes iguales. Michele Aguilar será socia y abogada titular: usará las "
        "herramientas, confirmará el criterio jurídico, firmará, aceptará o rechazará asuntos, y conservará veto "
        "sobre ética, conflictos y comunicaciones públicas del despacho. Ricardo De Biase será socio y responsable "
        "técnico: construirá y operará las herramientas según las prioridades que usted fije. No orientará clientes, "
        "no firmará, no litigará y no se anunciará como abogado.",
    )
    h2(doc, "2.2. Titularidad de la tecnología")
    body(
        doc,
        "DBX Solutions es una sociedad de titularidad exclusiva de Ricardo De Biase. En ella se concibió, se construyó "
        "y se mantiene la propiedad intelectual y la arquitectura de la inteligencia artificial. Esa titularidad no "
        "se cede ni se coparticipa por el hecho de constituir LexiaTek.",
    )
    h2(doc, "2.3. Datos por completar")
    note(doc, "Complete o corrija únicamente los campos entre corchetes. El nombre de las partes ya está fijado.")
    add_table(
        doc,
        ["Dato", "Valor"],
        [
            ["Nombre de la abogada titular", "Michele Aguilar"],
            ["Cédula de ciudadanía", "[          ]"],
            ["Tarjeta profesional", "[          ]"],
            ["Domicilio profesional", "[ciudad / dirección]"],
            ["Nombre de Ricardo De Biase", "Ricardo De Biase"],
            ["Nombre legal propuesto para LexiaTek", "[p. ej. LexiaTek S.A.S. — usted lo confirma]"],
            ["Domicilio de LexiaTek", "[ciudad de la Cámara de Comercio]"],
            ["Razón social y NIT de DBX Solutions", "[si ya está constituida; si no, indicarlo]"],
            ["Correos de notificación", "[correo de Michele Aguilar]  /  [correo de Ricardo De Biase]"],
        ],
    )

    h1(doc, "3. Propósito profesional")
    body(
        doc,
        "Un asunto jurídico —penal, civil, laboral, de tutela o de simple reclamación ante una entidad— exige tiempo, "
        "claridad y responsabilidad. El modelo tradicional de honorarios deja fuera a quienes más necesitan orientación. "
        "Escalar un despacho sin método pone en riesgo la calidad del trabajo y la tarjeta profesional. LexiaTek se "
        "propone atender más personas, con costos comprensibles y por etapas, sin relajar el rigor ni prometer resultados.",
    )
    body(
        doc,
        "Lo que se ofrece a usted no es un empleo. Es la titularidad profesional de un despacho que ya tiene marca, "
        "presencia pública y capacidad técnica, y que no puede existir de manera ética sin una abogada que decida y firme.",
    )

    h1(doc, "4. Estructura propuesta")
    body(
        doc,
        "Se propone constituir LexiaTek como el vehículo de la asociación: un despacho con dos socios, Michele Aguilar "
        "y Ricardo De Biase, titulares cada uno del cincuenta por ciento. LexiaTek contrata con el cliente, percibe "
        "el precio del servicio, cubre los costos de operación y reparte las utilidades netas.",
    )
    body(
        doc,
        "La forma societaria exacta —incluido el nombre legal, que hoy figura de manera provisional—, el registro ante "
        "Cámara de Comercio, el RUT y los demás actos de constitución los confirma usted, de modo que su tarjeta "
        "profesional y la participación de un no abogado en el capital del negocio queden alineadas con el régimen "
        "aplicable. Esta propuesta no sustituye ese examen.",
    )
    h2(doc, "4.1. Licencias recíprocas, vigentes solo mientras subsista la asociación")
    body(
        doc,
        "El esquema no consiste en ceder la inteligencia artificial ni en ceder de manera indefinida el ejercicio "
        "profesional. Consiste en dos autorizaciones recíprocas, de igual condición temporal:",
    )
    bullet(
        doc,
        "Licencia de la tecnología. La propiedad intelectual y la arquitectura de la inteligencia artificial "
        "pertenecen de manera exclusiva a DBX Solutions. DBX Solutions concede a Michele Aguilar y, a través de ella, "
        "a LexiaTek, una licencia de uso no exclusiva, limitada al objeto del despacho y vigente únicamente mientras "
        "subsista la asociación. No hay cesión ni copropiedad de esa tecnología.",
    )
    bullet(
        doc,
        "Autorización de los servicios profesionales. Michele Aguilar autoriza a LexiaTek a contar con sus servicios "
        "profesionales como abogada titular —orientación, revisión, firma y representación— únicamente mientras "
        "subsista la asociación. Al terminar la asociación, esa autorización cesa, sin perjuicio del periodo "
        "transitorio que usted confirme para la no abandono de clientes y la entrega ordenada de expedientes, "
        "conforme a sus deberes profesionales.",
    )
    body(
        doc,
        "Ambas autorizaciones nacen con la asociación y se extinguen con ella, salvo el periodo transitorio que usted "
        "determine. Ser socia al cincuenta por ciento de LexiaTek no convierte a Michele Aguilar en titular de la "
        "inteligencia artificial. Ser el único socio de DBX Solutions no autoriza a Ricardo De Biase a ejercer el "
        "derecho ni a decidir un escrito.",
    )

    h1(doc, "5. Distribución de funciones")
    add_table(
        doc,
        ["Materia", "Michele Aguilar", "Ricardo De Biase"],
        [
            ["Asuntos, firma y radicación", "Decisión exclusiva, como abogada titular", "No interviene"],
            ["Ética, conflictos y textos públicos", "Veto", "Ejecuta lo aprobado"],
            ["Propiedad intelectual de la IA", "Licenciataria mientras dure la asociación", "Titular exclusivo, a través de DBX Solutions"],
            ["Servicios profesionales en LexiaTek", "Los autoriza al despacho mientras dure la asociación", "No presta servicios jurídicos"],
            ["Prioridad de herramientas internas", "Define qué se construye primero", "Construye y opera lo priorizado"],
            ["Marca LexiaTek, contrataciones y gasto", "Decisión conjunta", "Decisión conjunta"],
            ["Catálogo de servicios y tarifas", "Acuerdo conjunto; veto ético del cobro", "Acuerdo conjunto; se elabora después"],
        ],
    )
    body(
        doc,
        "Si usted no puede atender un asunto, la suplencia solo podrá recaer en otro abogado, empleado o contratista "
        "del despacho. Ni Ricardo De Biase ni el sistema de inteligencia artificial pueden ocupar esa función.",
    )

    h1(doc, "6. Modelo económico")
    body(
        doc,
        "El régimen económico que se propone rige las utilidades de LexiaTek. DBX Solutions no se reparte.",
    )
    para(doc, "Orden de aplicación", size=11, bold=True, color=NAVY, space_after=6)
    bullet(
        doc,
        "Ingreso. El cliente paga a LexiaTek el servicio contratado. El catálogo y las tarifas se definirán "
        "después, de común acuerdo.",
    )
    bullet(
        doc,
        "Operación. Con cargo a ese ingreso se cubren, en primer término, los costos necesarios para mantener "
        "el despacho en funcionamiento: inteligencia artificial, hospedaje, sitio, herramientas y, cuando existan, "
        "remuneraciones de empleados o contratistas. Los servicios tecnológicos los prestará o facturará DBX Solutions, "
        "o terceros, como costo de operación, no como participación en el honorario profesional.",
    )
    bullet(
        doc,
        "Utilidad. El remanente se distribuye por partes iguales entre Michele Aguilar y Ricardo De Biase. "
        "No hay distribución si no hay utilidad. No se retiran fondos sobre déficit.",
    )
    body(
        doc,
        "Lo que LexiaTek pague a DBX Solutions es contraprestación por tecnología y operación, no la mitad de un "
        "honorario de abogado. Los gastos propios de DBX Solutions los decide su único socio. Los gastos de LexiaTek "
        "que superen un tope que ambos acuerden requerirán la firma de los dos. Si al inicio no hay ingresos "
        "suficientes, se acordará por escrito quién anticipa la operación y cómo se reembolsa.",
    )

    h1(doc, "7. Propiedad intelectual, autorización profesional y cese")
    body(
        doc,
        "La propiedad intelectual y la arquitectura de la inteligencia artificial —código, instrucciones de sistema, "
        "herramientas y desarrollos conexos— pertenecen a DBX Solutions. No hay cesión ni copropiedad a favor de "
        "LexiaTek ni de Michele Aguilar.",
    )
    body(
        doc,
        "La marca comercial LexiaTek, en cuanto identifica al despacho, pertenecerá al negocio de los dos socios. "
        "Los expedientes y los datos de los clientes son del despacho y de la responsabilidad profesional de la "
        "abogada titular. En ningún caso se utilizarán como conjunto de entrenamiento.",
    )
    body(
        doc,
        "Al terminar la asociación: (i) la inteligencia artificial permanece en DBX Solutions y cesa la licencia "
        "de uso a LexiaTek, salvo el periodo transitorio que usted confirme; (ii) cesa la autorización de sus "
        "servicios profesionales a LexiaTek, con el mismo periodo transitorio, a fin de no abandonar clientes y "
        "de ordenar la entrega de expedientes, en los términos que impongan sus deberes profesionales. Usted "
        "confirma la duración y las condiciones éticas de ese periodo.",
    )

    h1(doc, "8. Gobierno")
    body(
        doc,
        "En LexiaTek, las participaciones iguales no significan que el criterio jurídico se vote. El veto de Michele "
        "Aguilar sobre asuntos, firma, ética, conflictos y publicidad no se somete a mayoría. El veto técnico de "
        "Ricardo De Biase —seguridad, infraestructura y lo que no sea posible construir— tampoco se somete a mayoría.",
    )
    body(
        doc,
        "Las decisiones sobre marca, catálogo de servicios, contrataciones y gastos del despacho se adoptan de consuno. "
        "El empate en esas materias se resolverá por mediación o, en última instancia, por el derecho de uno de los "
        "socios a adquirir la participación del otro, en los términos que se fijen al constituir.",
    )

    h1(doc, "9. Personal futuro")
    body(
        doc,
        "Los abogados e ingenieros que se incorporen a LexiaTek lo harán como empleados o como contratistas. "
        "No adquirirán participación social, no percibirán el residual de los socios y no tendrán voto de dueños.",
    )
    body(
        doc,
        "Los contratos incluirán confidencialidad y la prohibición de solicitar o llevarse clientes del despacho. "
        "La contratación se acordará cuando exista caja suficiente.",
    )

    h1(doc, "10. Deberes éticos y profesionales")
    body(
        doc,
        "El secreto profesional, la verificación de conflictos, el consentimiento informado sobre el uso de "
        "herramientas tecnológicas, la Ley 1581 de 2012, las reglas de publicidad de la abogacía y el seguro de "
        "responsabilidad civil profesional son condiciones de operación. Si un borrador preparado con apoyo "
        "tecnológico se radica de manera indebida, la responsabilidad profesional recae en quien firma: la abogada "
        "titular. Por eso el veto jurídico no es negociable.",
    )
    body(
        doc,
        "La comunicación pública de LexiaTek no garantizará resultados. El sitio actual deberá ser auditado por "
        "usted, inclusive en su alcance temático, antes de atar su nombre al despacho. No se captará clientela "
        "de manera activa hasta que la titularidad profesional y el texto público estén formalizados.",
    )

    h1(doc, "11. Terminación")
    body(
        doc,
        "LexiaTek tendrá únicamente dos socios. Si uno se retira, el otro gozará de primera opción para adquirir "
        "esa participación. Los expedientes permanecerán sujetos a las reglas profesionales que usted indique. "
        "La propiedad intelectual de la inteligencia artificial permanecerá en DBX Solutions. Cesan las dos "
        "autorizaciones recíprocas, con el periodo transitorio ya indicado.",
    )
    body(
        doc,
        "El plazo de preaviso y la fórmula de valoración de la participación se fijarán en el instrumento constitutivo.",
    )

    h1(doc, "12. Estado actual de lo construido")
    add_table(
        doc,
        ["Elemento", "Estado", "Titularidad o efecto"],
        [
            ["Sitio lexiatek.com y páginas legales", "Construido", "Lo opera LexiaTek; lo construyó DBX Solutions. Usted audita el texto y el alcance."],
            ["Recepción de asuntos", "Construido", "Tecnología de DBX Solutions, licenciada al despacho mientras dure la asociación."],
            ["WhatsApp", "Pendiente", "Usted decide si abre ese canal y con qué tono."],
            ["Nombre legal", "Provisional", "Usted lo define al constituir."],
            ["Herramientas internas de apoyo", "Construidas o diseñadas por DBX Solutions", "IP de DBX Solutions. Usted aprueba cada salida jurídica."],
            ["Abogada titular", "Por formalizar", "Michele Aguilar. Sin usted no hay despacho."],
        ],
    )

    h1(doc, "13. Qué sigue, si usted está de acuerdo en lo sustancial")
    body(
        doc,
        "Esta propuesta no abre el despacho por sí sola. En Colombia el orden práctico es: primero alinear el texto; "
        "después firmar los contratos entre las partes; después constituir las sociedades que falten; después cumplir "
        "el régimen profesional, tributario y de datos; y solo entonces atender al público. Usted confirma cada paso "
        "jurídico. Lo que sigue es una hoja de ruta de trabajo, no un dictamen.",
    )
    h2(doc, "Etapa A. Cerrar este borrador")
    bullet(doc, "Usted revisa este documento con control de cambios y ajusta el lenguaje hasta que le represente.")
    bullet(doc, "Ambos congelan el cuadro de términos del numeral 15.")
    bullet(doc, "Usted confirma el tipo societario de LexiaTek y si un no abogado puede figurar como socio de esa persona jurídica según el objeto que se registre.")
    bullet(doc, "Se verifica el estado de DBX Solutions: matrícula, NIT y objeto social suficiente para desarrollar y licenciar software.")
    h2(doc, "Etapa B. Contratos entre las partes")
    bullet(doc, "Pacto de socios de LexiaTek: participaciones 50/50, vetos, cascada de utilidades, empate, personal sin participación social, salida y primera opción de compra.")
    bullet(doc, "Contrato de licencia de la propiedad intelectual y de la inteligencia artificial: DBX Solutions a Michele Aguilar y, a través de ella, a LexiaTek, vigente únicamente mientras subsista la asociación, con periodo transitorio que usted confirme.")
    bullet(doc, "Autorización o licencia de servicios profesionales: Michele Aguilar a LexiaTek, vigente únicamente mientras subsista la asociación, con el mismo periodo transitorio para no abandonar clientes.")
    bullet(doc, "Cláusulas de confidencialidad, secreto profesional y tratamiento de datos, si no quedan íntegras en los instrumentos anteriores.")
    h2(doc, "Etapa C. DBX Solutions")
    bullet(doc, "Si ya existe: certificado de existencia y representación legal, RUT, objeto social y cuenta bancaria a su nombre.")
    bullet(doc, "Si aún no existe: constituirla a nombre exclusivo de Ricardo De Biase, para que el licenciante de la tecnología sea identificable.")
    h2(doc, "Etapa D. Constitución de LexiaTek")
    bullet(doc, "Consulta de homonimia en el RUES y reserva del nombre que usted apruebe.")
    bullet(doc, "Estatutos acordes con el pacto de socios.")
    bullet(doc, "Documento de constitución. En S.A.S. suele bastar documento privado; usted indica si se requiere escritura pública.")
    bullet(doc, "Registro en la Cámara de Comercio del domicilio social.")
    bullet(doc, "RUT ante la DIAN.")
    bullet(doc, "Cuenta bancaria a nombre de LexiaTek. Los ingresos del despacho no se mezclan con cuentas personales ni con las de DBX Solutions.")
    bullet(doc, "Contador público. Facturación electrónica cuando sea obligatorio.")
    h2(doc, "Etapa E. Encaje profesional y de cumplimiento")
    bullet(doc, "Tarjeta profesional vigente y domicilio profesional conocido.")
    bullet(doc, "Seguro de responsabilidad civil profesional.")
    bullet(doc, "Régimen de poderes y de firma: siempre usted u otro abogado del despacho.")
    bullet(doc, "Política de tratamiento de datos y, si corresponde, Registro Nacional de Bases de Datos.")
    bullet(doc, "Auditoría del sitio y de la publicidad, incluido el alcance temático, antes de captar clientela.")
    h2(doc, "Etapa F. Operación")
    bullet(doc, "Apertura de WhatsApp, si usted lo autoriza.")
    bullet(doc, "Catálogo de servicios y tarifas, de común acuerdo, con su veto ético.")
    bullet(doc, "Atención de asuntos a nombre del despacho ya constituido.")
    h2(doc, "Quién hace qué")
    add_table(
        doc,
        ["Paso", "Michele Aguilar", "Ricardo De Biase"],
        [
            ["Lenguaje de esta propuesta", "Corrige y aprueba", "Ajusta el borrador según sus cambios"],
            ["Tipo societario de LexiaTek", "Define y confirma", "Aporta lo técnico; no dictamina"],
            ["Pacto de socios y licencias recíprocas", "Revisa y firma", "Revisa y firma"],
            ["Constitución de LexiaTek", "Revisa estatutos y firma", "Gestiona logística y costos de registro, si así se acuerda"],
            ["DBX Solutions", "No interviene como socia", "Constituye o actualiza; es el único socio"],
            ["Sitio, datos y publicidad", "Veta y aprueba el texto", "Ejecuta los cambios técnicos"],
            ["Tarifas al público", "Veto ético; acuerdo conjunto", "Acuerdo conjunto"],
            ["Captación de clientes", "Autoriza el momento", "No capta a nombre propio como abogado"],
        ],
    )

    h1(doc, "14. Documentos a suscribir")
    bullet(doc, "Pacto de socios de LexiaTek.")
    bullet(doc, "Estatutos y actos de constitución de LexiaTek, en la forma que usted confirme.")
    bullet(doc, "Contrato de licencia de la inteligencia artificial: DBX Solutions a Michele Aguilar y a LexiaTek, mientras subsista la asociación.")
    bullet(doc, "Autorización de servicios profesionales: Michele Aguilar a LexiaTek, mientras subsista la asociación.")
    body(
        doc,
        "El catálogo de servicios al público y sus tarifas no se suscriben con este borrador. Se elaborarán después, de común acuerdo.",
    )

    h1(doc, "15. Cuadro de términos")
    body(doc, "Síntesis para su anotación. No reemplaza los instrumentos del numeral 14.")
    add_table(
        doc,
        ["Término", "Propuesta"],
        [
            ["LexiaTek", "Despacho. Dos socios: Michele Aguilar y Ricardo De Biase, 50 % cada uno."],
            ["Michele Aguilar en LexiaTek", "Socia y abogada titular."],
            ["Ricardo De Biase en LexiaTek", "Socio y responsable técnico. No ejerce el derecho."],
            ["DBX Solutions", "Sociedad tecnológica. Socio único: Ricardo De Biase."],
            ["IP de la IA", "Exclusiva de DBX Solutions. Licencia a Michele Aguilar y a LexiaTek solo mientras dure la asociación."],
            ["Servicios profesionales", "Michele Aguilar los autoriza a LexiaTek solo mientras dure la asociación."],
            ["Cese", "Al terminar la asociación cesan ambas autorizaciones, con periodo transitorio que usted confirme."],
            ["Alcance del despacho", "Penal, civil, laboral, acoso sexual, PQRS, tutelas y trámites frecuentes de la ciudadanía."],
            ["Ingresos", "Los percibe LexiaTek."],
            ["Operación", "Tecnología, hospedaje, sitio y, luego, nómina. Se cubre antes de repartir."],
            ["Utilidades de LexiaTek", "50 % / 50 %. Sin utilidad, no hay reparto."],
            ["Expedientes", "Del despacho. Responsabilidad profesional de Michele Aguilar."],
            ["Personal futuro", "Empleados o contratistas. Sin participación social."],
            ["Tarifas al público", "Se definen después. Usted veta el cobro éticamente inadmisible."],
            ["Próximo paso", "Usted corrige este borrador. Si hay acuerdo sustancial, se redactan los contratos y se constituye."],
        ],
    )

    h1(doc, "16. Solicitud")
    body(
        doc,
        "Se solicita a Michele Aguilar que examine esta propuesta y manifieste si está de acuerdo en formalizar "
        "LexiaTek como despacho de dos socios, en partes iguales; en que la inteligencia artificial permanezca "
        "en DBX Solutions y se licencie al despacho solo mientras subsista la asociación; y en que sus servicios "
        "profesionales se autoricen a LexiaTek bajo la misma condición temporal.",
    )
    para(
        doc,
        "LexiaTek de ambos. La inteligencia artificial, de DBX Solutions, licenciada al despacho mientras dure "
        "la asociación. Los servicios profesionales de Michele Aguilar, autorizados al despacho mientras dure "
        "la asociación. Queda a su consideración formalizar la asociación en esos términos.",
        size=12, italic=True, color=NAVY, align="center", space_before=16, space_after=20,
    )
    para(doc, "Ricardo De Biase", size=12, bold=True, color=NAVY, align="center", space_after=0)
    para(doc, "DBX Solutions  ·  " + FECHA, size=10, color=MUTED, align="center", space_after=12)
    para(
        doc,
        "Documento de trabajo (borrador). No constituye oferta irrevocable, contrato ni dictamen. "
        "Cualquier instrumento posterior requerirá la revisión y la firma de Michele Aguilar y de Ricardo De Biase.",
        size=9, italic=True, color=MUTED, align="center",
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    build()
