#!/usr/bin/env python3
"""Deck ampliado de onboarding LexiaTek: agentes, skills, herramientas y tareas."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


# LexiaTek visual system
NAVY = RGBColor(0x0F, 0x17, 0x2A)
BRAND = RGBColor(0x1E, 0x3A, 0x8A)
BLUE = RGBColor(0x25, 0x63, 0xEB)
CYAN = RGBColor(0x06, 0xB6, 0xD4)
MIST = RGBColor(0xE6, 0xEE, 0xF7)
SOFT = RGBColor(0xDC, 0xE7, 0xF2)
SURFACE = RGBColor(0xF4, 0xF8, 0xFC)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
HEADING = RGBColor(0x0F, 0x17, 0x2A)
BODY = RGBColor(0x47, 0x55, 0x69)
MUTED = RGBColor(0x64, 0x74, 0x8B)
BORDER = RGBColor(0xD5, 0xE0, 0xEB)
GREEN = RGBColor(0x0F, 0x76, 0x6E)
AMBER = RGBColor(0xB4, 0x53, 0x09)

FONT_H = "Poppins"
FONT_B = "Lato"
W = Inches(13.333)
H = Inches(7.5)
OUT = Path(__file__).resolve().parent / "LexiaTek_Onboarding_Despacho_Virtual.pptx"


AGENTS = [
    {
        "name": "Coordinador del Caso",
        "role": "Dirección operativa del expediente y voz única del despacho.",
        "why": "Evita respuestas dispersas, identifica urgencias y no permite avanzar con datos críticos incompletos.",
        "does": [
            "Recibe la solicitud en lenguaje natural.",
            "Clasifica tarea, etapa y prioridad.",
            "Pide datos o documentos faltantes.",
            "Asigna el trabajo al área correcta.",
            "Consolida la respuesta para la abogada.",
        ],
        "skills": [
            "Clasificar tarea y etapa",
            "Detectar urgencia penal",
            "Gestionar faltantes",
            "Marcar pendientes de verificación",
            "Actualizar tareas y responsable",
        ],
        "io": "Entrada: solicitud + expediente  →  Salida: ruta de trabajo, faltantes y siguiente paso",
        "limit": "No sustituye el análisis especializado ni la aprobación final.",
    },
    {
        "name": "Cronología y Hechos",
        "role": "Reconstrucción factual ordenada, trazable y verificable.",
        "why": "Una inconsistencia en fechas, actores o fuentes debilita tipicidad, escritos y credibilidad.",
        "does": [
            "Extrae hechos relevantes.",
            "Ordena eventos por fecha y actor.",
            "Separa confirmado, narrado e inferido.",
            "Detecta contradicciones y vacíos.",
            "Construye matriz hecho–fuente.",
        ],
        "skills": [
            "Construir cronología penal",
            "Extraer hechos relevantes",
            "Detectar contradicciones",
            "Detectar vacíos factuales",
            "Crear matriz hecho–fuente",
        ],
        "io": "Entrada: relato + documentos  →  Salida: línea de tiempo, fuentes y preguntas faltantes",
        "limit": "No realiza la calificación penal definitiva.",
    },
    {
        "name": "Tipicidad y Responsabilidad",
        "role": "Análisis sustantivo preliminar de posibles conductas punibles.",
        "why": "Evita orientar actuaciones sobre hipótesis típicas débiles o sin soporte fáctico y probatorio.",
        "does": [
            "Descompone elementos del tipo penal.",
            "Mapea hecho, elemento y prueba.",
            "Analiza autoría, participación y dolo/culpa.",
            "Identifica agravantes y atenuantes.",
            "Señala riesgos de atipicidad.",
        ],
        "skills": [
            "Descomponer elementos del tipo",
            "Identificar conductas preliminares",
            "Mapear tipo–hecho–prueba",
            "Analizar autoría y dolo/culpa",
            "Detectar riesgos de atipicidad",
        ],
        "io": "Entrada: hechos + evidencia  →  Salida: hipótesis preliminar y brechas por probar",
        "limit": "No decide imputación, acusación ni estrategia final.",
    },
    {
        "name": "Ruta Procesal Ley 906",
        "role": "Ubicación de etapa, oportunidad y actuación posible para la víctima.",
        "why": "Reduce riesgos de extemporaneidad, improcedencia y solicitudes dirigidas al actor equivocado.",
        "does": [
            "Identifica etapa procesal Ley 906.",
            "Mapea actuaciones posibles de la víctima.",
            "Evalúa oportunidad y términos.",
            "Detecta riesgos procesales.",
            "Propone una ruta recomendada.",
        ],
        "skills": [
            "Identificar etapa Ley 906",
            "Mapear actuaciones de la víctima",
            "Evaluar oportunidad procesal",
            "Control preliminar de términos",
            "Crear ruta procesal",
        ],
        "io": "Entrada: estado + última actuación  →  Salida: etapa, opciones, riesgos y próximo paso",
        "limit": "No sustituye el seguimiento operativo diario del radicado.",
    },
    {
        "name": "Representación de Víctimas",
        "role": "Estrategia centrada en derechos, intereses y no revictimización.",
        "why": "Una estrategia puede ser técnicamente correcta y aun así ignorar el objetivo real o el riesgo de la víctima.",
        "does": [
            "Identifica intereses y derechos.",
            "Construye teoría del caso de víctima.",
            "Evalúa daño y afectación.",
            "Aplica enfoque diferencial.",
            "Prioriza objetivos sin revictimizar.",
        ],
        "skills": [
            "Identificar intereses de la víctima",
            "Construir teoría del caso",
            "Analizar derechos de víctima",
            "Detectar riesgo de revictimización",
            "Priorizar objetivos",
        ],
        "io": "Entrada: hechos + contexto de la víctima  →  Salida: objetivos y enfoque de representación",
        "limit": "No promete resultados ni sustituye decisiones sensibles del despacho.",
    },
    {
        "name": "Evidencia y Pruebas",
        "role": "Gestión probatoria: inventario, soporte, brechas y plan de recaudo.",
        "why": "Sin trazabilidad y suficiencia probatoria, la teoría del caso pierde fuerza en escritos y audiencias.",
        "does": [
            "Inventaría y clasifica evidencia.",
            "Construye matriz hecho–prueba.",
            "Detecta brechas probatorias.",
            "Propone plan de recaudo.",
            "Marca riesgos de evidencia digital y custodia.",
        ],
        "skills": [
            "Inventariar evidencia",
            "Clasificar tipo de prueba",
            "Construir matriz hecho–prueba",
            "Detectar brechas probatorias",
            "Crear plan de recaudo",
        ],
        "io": "Entrada: soportes + hechos  →  Salida: inventario, suficiencia y acciones de recaudo",
        "limit": "No reemplaza peritaje forense ni certifica cadena de custodia.",
    },
    {
        "name": "Audiencias Penales",
        "role": "Preparación táctica de la intervención de la víctima en audiencia.",
        "why": "Reduce improvisación, omisiones y preguntas que puedan revictimizar.",
        "does": [
            "Define objetivo jurídico y táctico.",
            "Prepara guion y solicitudes orales.",
            "Formula preguntas y contraargumentos.",
            "Simula escenarios y riesgos.",
            "Entrega checklist previo.",
        ],
        "skills": [
            "Identificar objetivo de audiencia",
            "Preparar guion oral",
            "Preparar solicitudes y preguntas",
            "Simular escenarios",
            "Detectar riesgos de audiencia",
        ],
        "io": "Entrada: tipo de audiencia + objetivo + prueba  →  Salida: guion, solicitudes y checklist",
        "limit": "No reemplaza la intervención oral de la abogada en estrados.",
    },
    {
        "name": "Redacción de Documentos",
        "role": "Producción de borradores jurídicos estructurados y revisables.",
        "why": "Acelera el primer borrador sin sacrificar estructura, soporte ni trazabilidad de pendientes.",
        "does": [
            "Redacta memoriales e impulsos.",
            "Prepara derechos de petición penales.",
            "Estructura hechos, fundamentos y solicitudes.",
            "Marca citas, anexos o radicados por verificar.",
            "Ajusta tono jurídico y separación hecho/inferencia.",
        ],
        "skills": [
            "Redactar memorial penal",
            "Redactar impulso procesal",
            "Redactar derecho de petición",
            "Estructurar hechos y solicitudes",
            "Verificar citas normativas",
        ],
        "io": "Entrada: plan aprobado + insumos verificados  →  Salida: borrador para revisión y firma",
        "limit": "Alto riesgo: nunca radica ni sustituye el criterio de firma.",
    },
    {
        "name": "Seguimiento Procesal",
        "role": "Continuidad operativa del expediente, actuaciones, audiencias y términos.",
        "why": "Evita pérdida de oportunidad por inactividad, falta de control o ausencia de alertas.",
        "does": [
            "Registra actuaciones y radicados.",
            "Controla audiencias y documentos enviados.",
            "Genera alertas de términos.",
            "Detecta inactividad procesal.",
            "Produce reporte de estado y tareas.",
        ],
        "skills": [
            "Monitorear radicado",
            "Registrar actuación",
            "Generar alertas de términos",
            "Detectar inactividad",
            "Crear reporte de estado",
        ],
        "io": "Entrada: radicado + actuaciones  →  Salida: estado, alertas, pendientes y responsable",
        "limit": "No reemplaza el análisis jurídico estratégico de la actuación.",
    },
    {
        "name": "Control de Calidad Jurídica",
        "role": "Puerta de control antes de cualquier salida externa.",
        "why": "Disminuye riesgo de hechos sin soporte, citas falsas, incoherencia, datos sensibles y revictimización.",
        "does": [
            "Verifica soporte fáctico y citas.",
            "Detecta alucinaciones legales.",
            "Revisa coherencia de estrategia y prueba.",
            "Controla confidencialidad y tono.",
            "Clasifica: aprobable, requiere cambios o rechazar.",
        ],
        "skills": [
            "Verificar hechos y citas",
            "Detectar alucinaciones legales",
            "Controlar confidencialidad",
            "Revisar coherencia estratégica",
            "Clasificar aprobación jurídica",
        ],
        "io": "Entrada: análisis o borrador  →  Salida: hallazgos, cambios exigidos y dictamen técnico",
        "limit": "No aprueba automáticamente ni reemplaza la responsabilidad profesional.",
    },
]


def add_run(paragraph, text, size=16, bold=False, color=BODY, font=FONT_B):
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return run


def text(slide, x, y, w, h, value, size=16, bold=False, color=BODY,
         align=PP_ALIGN.LEFT, font=FONT_B, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(x, y, w, h)
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = anchor
    p = frame.paragraphs[0]
    p.alignment = align
    add_run(p, value, size, bold, color, font)
    return box


def rich_text(slide, x, y, w, h, blocks, size=14, color=BODY, gap=4):
    box = slide.shapes.add_textbox(x, y, w, h)
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    for i, block in enumerate(blocks):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.space_after = Pt(gap)
        if isinstance(block, tuple):
            label, value = block
            add_run(p, label, size, True, HEADING, FONT_H)
            add_run(p, value, size, False, color, FONT_B)
        else:
            add_run(p, block, size, False, color, FONT_B)
    return box


def fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def rect(slide, x, y, w, h, color, radius=True, line=None):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, x, y, w, h)
    fill(shape, color)
    if line:
        shape.line.color.rgb = line
        shape.line.width = Pt(1)
    if radius:
        try:
            shape.adjustments[0] = 0.1
        except Exception:
            pass
    return shape


def set_bg(slide, color=MIST):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def footer(slide, number, section="Onboarding del despacho"):
    text(slide, Inches(0.7), Inches(7.08), Inches(5), Inches(0.2),
         f"LexiaTek · {section}", 9, False, MUTED)
    text(slide, Inches(12.2), Inches(7.08), Inches(0.45), Inches(0.2),
         str(number), 9, True, MUTED, PP_ALIGN.RIGHT)


def header(slide, title, subtitle, number, section="Onboarding del despacho"):
    rect(slide, Inches(0), Inches(0), W, Inches(0.08), CYAN, radius=False)
    text(slide, Inches(0.7), Inches(0.24), Inches(6), Inches(0.25),
         section.upper(), 10, True, CYAN, font=FONT_H)
    text(slide, Inches(0.7), Inches(0.58), Inches(12), Inches(0.52),
         title, 26, True, HEADING, font=FONT_H)
    text(slide, Inches(0.7), Inches(1.15), Inches(11.8), Inches(0.4),
         subtitle, 13, False, MUTED)
    footer(slide, number, section)


def card(slide, x, y, w, h, title_value, body, accent=BRAND, body_size=13):
    rect(slide, x, y, w, h, WHITE, line=BORDER)
    rect(slide, x, y, Inches(0.1), h, accent, radius=False)
    text(slide, x + Inches(0.25), y + Inches(0.18), w - Inches(0.4), Inches(0.35),
         title_value, 14, True, accent, font=FONT_H)
    text(slide, x + Inches(0.25), y + Inches(0.58), w - Inches(0.45), h - Inches(0.72),
         body, body_size, False, BODY)


def pill(slide, x, y, w, label, color=BRAND):
    rect(slide, x, y, w, Inches(0.38), SURFACE, line=BORDER)
    text(slide, x + Inches(0.08), y + Inches(0.08), w - Inches(0.16), Inches(0.2),
         label, 10, True, color, PP_ALIGN.CENTER, FONT_H)


def arrow(slide, x, y, w=Inches(0.35)):
    shape = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x, y, w, Inches(0.22))
    fill(shape, CYAN)
    return shape


def blank(prs, bg=MIST):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, bg)
    return slide


def cover(prs):
    slide = blank(prs, NAVY)
    rect(slide, Inches(0), Inches(0), Inches(0.16), H, CYAN, radius=False)
    text(slide, Inches(0.9), Inches(1.15), Inches(6.8), Inches(0.35),
         "LEXIATEK · ONBOARDING", 13, True, CYAN, font=FONT_H)
    text(slide, Inches(0.9), Inches(1.75), Inches(7.0), Inches(1.5),
         "Cómo funciona el\ndespacho virtual", 38, True, WHITE, font=FONT_H)
    text(slide, Inches(0.9), Inches(3.55), Inches(6.7), Inches(0.8),
         "Agentes dedicados · skills procesales · herramientas · tareas de inicio a cierre",
         17, False, SOFT)
    rect(slide, Inches(8.25), Inches(1.0), Inches(4.25), Inches(5.35), BRAND)
    text(slide, Inches(8.62), Inches(1.55), Inches(3.5), Inches(0.3),
         "EL MODELO", 11, True, CYAN, font=FONT_H)
    text(slide, Inches(8.62), Inches(2.05), Inches(3.5), Inches(1.45),
         "1 Coordinador\n9 especialistas\n1 decisión humana", 24, True, WHITE, font=FONT_H)
    text(slide, Inches(8.62), Inches(4.0), Inches(3.45), Inches(1.3),
         "La IA propone.\nLa abogada revisa,\najusta y aprueba.", 16, False, SOFT)
    text(slide, Inches(0.9), Inches(6.65), Inches(6), Inches(0.25),
         "Penal–víctimas · Colombia · Ley 906 de 2004", 11, False, MUTED)


def learning(prs, n):
    slide = blank(prs)
    header(slide, "Qué debe poder hacer al terminar", "Una guía práctica, no una clase de ingeniería.", n)
    goals = [
        ("Entender", "por qué existen agentes dedicados y cómo se reparten el trabajo."),
        ("Distinguir", "un agente, un skill y una herramienta."),
        ("Pedir", "una tarea en lenguaje natural con los insumos mínimos."),
        ("Supervisar", "qué agente responde, qué produce y dónde termina su responsabilidad."),
        ("Cerrar", "una tarea con revisión, aprobación, firma y seguimiento."),
    ]
    for i, (verb, body) in enumerate(goals):
        y = Inches(1.85) + i * Inches(0.92)
        rect(slide, Inches(0.9), y, Inches(0.58), Inches(0.58), BRAND)
        text(slide, Inches(0.9), y + Inches(0.13), Inches(0.58), Inches(0.28),
             str(i + 1), 16, True, WHITE, PP_ALIGN.CENTER, FONT_H)
        rich_text(slide, Inches(1.75), y + Inches(0.05), Inches(10.5), Inches(0.55),
                  [(verb + ": ", body)], 16, BODY, 0)


def dedicated_agents(prs, n):
    slide = blank(prs)
    header(slide, "El concepto: agentes dedicados a funciones", "Cada agente se comporta como un área del despacho con un mandato limitado.", n, "Conceptos")
    card(slide, Inches(0.7), Inches(1.85), Inches(3.75), Inches(4.6),
         "Agente generalista", "Intenta responder de todo.\n\n• Mezcla funciones\n• Difícil saber quién hizo qué\n• Más riesgo de omisiones\n• Control de calidad difuso", MUTED, 15)
    card(slide, Inches(4.8), Inches(1.85), Inches(3.75), Inches(4.6),
         "Agente dedicado", "Tiene una responsabilidad concreta.\n\n• Entradas definidas\n• Procedimiento definido\n• Salida esperada\n• Límite explícito\n• Trabajo auditable", BLUE, 15)
    card(slide, Inches(8.9), Inches(1.85), Inches(3.75), Inches(4.6),
         "Despacho coordinado", "Los agentes no compiten ni hablan todos a la vez.\n\nEl Coordinador asigna; los especialistas producen; Calidad revisa; la abogada decide.", CYAN, 15)


def agent_skill_tool(prs, n):
    slide = blank(prs)
    header(slide, "Agente, skill y herramienta no son lo mismo", "Tres capas que trabajan juntas para producir una salida jurídica controlada.", n, "Conceptos")
    layers = [
        ("AGENTE = QUIÉN", "Rol con mandato y responsabilidad.\n\nEj.: Ruta Procesal Ley 906 decide qué análisis procesal debe realizar.", BRAND),
        ("SKILL = CÓMO", "Playbook procedimental con propósito, entradas, pasos, salida y guardrails.\n\nEj.: identificar etapa procesal.", BLUE),
        ("HERRAMIENTA = CON QUÉ", "Función para consultar expediente o conocimiento autorizado.\n\nEj.: buscar en expediente o leer normas clave.", GREEN),
    ]
    for i, (title_value, body, color) in enumerate(layers):
        x = Inches(0.7) + i * Inches(4.15)
        card(slide, x, Inches(1.9), Inches(3.85), Inches(3.95), title_value, body, color, 14)
        if i < 2:
            arrow(slide, x + Inches(3.92), Inches(3.72), Inches(0.25))
    rect(slide, Inches(0.7), Inches(6.15), Inches(12), Inches(0.55), NAVY)
    text(slide, Inches(0.95), Inches(6.29), Inches(11.5), Inches(0.25),
         "Fórmula: el agente selecciona un skill, el skill usa herramientas y produce una salida revisable.",
         14, True, WHITE, PP_ALIGN.CENTER, FONT_H)


def skill_anatomy(prs, n):
    slide = blank(prs)
    header(slide, "Qué contiene un skill procesal", "Un skill convierte conocimiento jurídico en una tarea repetible y verificable.", n, "Conceptos")
    pieces = [
        ("1. Propósito", "Qué problema resuelve."),
        ("2. Entradas", "Hechos, documentos y contexto mínimos."),
        ("3. Pasos", "Secuencia concreta de análisis."),
        ("4. Salida", "Producto prometido y formato."),
        ("5. Guardrails", "Qué no inventar y cuándo detenerse."),
        ("6. Handoff", "Quién recibe el resultado después."),
    ]
    for i, (title_value, body) in enumerate(pieces):
        row, col = divmod(i, 3)
        x = Inches(0.7) + col * Inches(4.15)
        y = Inches(1.85) + row * Inches(2.05)
        card(slide, x, y, Inches(3.85), Inches(1.78), title_value, body, CYAN if i in (2, 4) else BRAND, 14)
    text(slide, Inches(0.9), Inches(6.25), Inches(11.6), Inches(0.45),
         "Ejemplo: identificar etapa Ley 906 → recibe actuaciones y fechas → contrasta expediente y playbook → entrega etapa, evidencia e incertidumbres.",
         13, False, BODY, PP_ALIGN.CENTER)


def tools_slide(prs, n):
    slide = blank(prs)
    header(slide, "Herramientas disponibles hoy", "Accesos implementados; no deben confundirse con capacidades futuras.", n, "Conceptos")
    tools = [
        ("Buscar en expediente", "Recupera hechos, documentos y actuaciones de la sesión vinculada."),
        ("Buscar en conocimiento", "Consulta la base jurídica y operativa autorizada."),
        ("Leer área de derecho", "Carga el marco temático relevante para especialistas."),
        ("Leer playbook del proceso", "Consulta etapas y flujo canónico de la Ley 906."),
        ("Leer normas clave", "Recupera criterios y normas penales de referencia."),
        ("Listar áreas de derecho", "Permite ubicar el área aplicable antes de profundizar."),
    ]
    for i, (name, body) in enumerate(tools):
        row, col = divmod(i, 2)
        x = Inches(0.7) + col * Inches(6.2)
        y = Inches(1.78) + row * Inches(1.55)
        card(slide, x, y, Inches(5.95), Inches(1.35), name, body, GREEN, 12)
    rect(slide, Inches(0.7), Inches(6.48), Inches(12), Inches(0.38), SOFT)
    text(slide, Inches(0.85), Inches(6.57), Inches(11.7), Inches(0.2),
         "Regla: si una herramienta o dato no existe, el sistema marca [PENDIENTE DE VERIFICAR]; no lo inventa.",
         11, True, AMBER, PP_ALIGN.CENTER, FONT_H)


def despacho(prs, n):
    slide = blank(prs)
    header(slide, "Qué es el despacho virtual LexiaTek", "Una organización de trabajo penal–víctimas, no un chatbot que reemplaza a la abogada.", n, "El despacho")
    rect(slide, Inches(0.7), Inches(1.9), Inches(12), Inches(1.0), NAVY)
    text(slide, Inches(1.0), Inches(2.12), Inches(11.4), Inches(0.55),
         "Un sistema que recibe un objetivo jurídico, divide el trabajo por funciones, aplica procedimientos especializados y devuelve productos para decisión humana.",
         18, True, WHITE, PP_ALIGN.CENTER, FONT_H)
    columns = [
        ("Alcance", "Colombia · penal con enfoque en representación de víctimas · Ley 906."),
        ("Cobertura", "Hechos, tipicidad, ruta, víctimas, prueba, audiencias, escritos, seguimiento y calidad."),
        ("Límite profesional", "No firma, no radica, no promete resultados y no sustituye criterio profesional."),
    ]
    for i, (title_value, body) in enumerate(columns):
        card(slide, Inches(0.7) + i * Inches(4.15), Inches(3.35), Inches(3.85), Inches(2.25),
             title_value, body, CYAN if i == 2 else BRAND, 14)
    text(slide, Inches(0.85), Inches(6.15), Inches(11.7), Inches(0.35),
         "Fuera del producto: tutela / vía constitucional y defensa del investigado.", 13, True, AMBER, PP_ALIGN.CENTER, FONT_H)


def architecture(prs, n):
    slide = blank(prs)
    header(slide, "Cómo se organiza el despacho", "Una sola voz hacia la abogada; trabajo modular y auditable por detrás.", n, "El despacho")
    boxes = [
        ("1. Abogada", "Define objetivo\ny aporta contexto", BRAND),
        ("2. Coordinador", "Clasifica, completa\ny asigna", NAVY),
        ("3. Especialistas", "Analizan o\nproducen", BLUE),
        ("4. Calidad", "Revisa soporte,\nriesgo y tono", GREEN),
        ("5. Abogada", "Ajusta, aprueba,\nfirma o rechaza", BRAND),
    ]
    x0, y = Inches(0.55), Inches(2.45)
    bw, gap = Inches(2.05), Inches(0.42)
    for i, (title_value, body, color) in enumerate(boxes):
        x = x0 + i * (bw + gap)
        rect(slide, x, y, bw, Inches(2.25), color)
        text(slide, x + Inches(0.12), y + Inches(0.3), bw - Inches(0.24), Inches(0.4),
             title_value, 14, True, WHITE, PP_ALIGN.CENTER, FONT_H)
        text(slide, x + Inches(0.12), y + Inches(0.92), bw - Inches(0.24), Inches(0.9),
             body, 13, False, SOFT, PP_ALIGN.CENTER)
        if i < len(boxes) - 1:
            arrow(slide, x + bw + Inches(0.08), y + Inches(1.0), Inches(0.25))
    text(slide, Inches(1.0), Inches(5.35), Inches(11.3), Inches(0.7),
         "Los especialistas no realizan handoffs entre sí ni hablan como personajes distintos. El Coordinador mantiene contexto, trazabilidad y voz única.",
         15, False, BODY, PP_ALIGN.CENTER)


def task_lifecycle(prs, n):
    slide = blank(prs)
    header(slide, "El ciclo de vida de cualquier tarea", "Saber dónde comienza y dónde termina evita expectativas incorrectas.", n, "El despacho")
    steps = [
        ("1", "Objetivo", "Qué necesita decidir o producir la abogada."),
        ("2", "Insumos", "Hechos, documentos, etapa y fecha relevante."),
        ("3", "Plan", "Agentes, skills, dependencias y salida esperada."),
        ("4", "Ejecución", "Consulta expediente/conocimiento y produce módulos."),
        ("5", "Control", "Calidad revisa soporte, coherencia, tono y riesgos."),
        ("6", "Cierre", "La abogada aprueba, firma y crea seguimiento."),
    ]
    for i, (num, title_value, body) in enumerate(steps):
        x = Inches(0.45) + i * Inches(2.13)
        rect(slide, x, Inches(2.15), Inches(1.85), Inches(3.2), WHITE, line=BORDER)
        rect(slide, x, Inches(2.15), Inches(1.85), Inches(0.55), NAVY if i in (0, 5) else BRAND)
        text(slide, x + Inches(0.08), Inches(2.28), Inches(0.35), Inches(0.25), num, 12, True, CYAN, font=FONT_H)
        text(slide, x + Inches(0.4), Inches(2.28), Inches(1.3), Inches(0.25), title_value, 12, True, WHITE, font=FONT_H)
        text(slide, x + Inches(0.15), Inches(3.0), Inches(1.55), Inches(1.8), body, 12, False, BODY, PP_ALIGN.CENTER)
        if i < 5:
            arrow(slide, x + Inches(1.9), Inches(3.55), Inches(0.2))
    rect(slide, Inches(0.85), Inches(5.85), Inches(11.65), Inches(0.62), SOFT)
    text(slide, Inches(1.0), Inches(6.02), Inches(11.35), Inches(0.25),
         "Terminar una tarea no siempre significa radicar: puede terminar en análisis, checklist, borrador, decisión o alerta.",
         13, True, BRAND, PP_ALIGN.CENTER, FONT_H)


def agents_map(prs, n):
    slide = blank(prs)
    header(slide, "Los 10 roles del despacho", "1 Coordinador + 8 áreas productoras + 1 puerta de calidad.", n, "Los agentes")
    rect(slide, Inches(4.55), Inches(1.78), Inches(4.25), Inches(0.72), NAVY)
    text(slide, Inches(4.7), Inches(2.0), Inches(3.95), Inches(0.28),
         "Coordinador del Caso", 15, True, WHITE, PP_ALIGN.CENTER, FONT_H)
    groups = [
        ("Base factual", ["Cronología y Hechos", "Evidencia y Pruebas"]),
        ("Análisis jurídico", ["Tipicidad y Responsabilidad", "Ruta Procesal Ley 906", "Representación de Víctimas"]),
        ("Producción", ["Audiencias Penales", "Redacción de Documentos", "Seguimiento Procesal"]),
    ]
    for gi, (label, names) in enumerate(groups):
        x = Inches(0.65) + gi * Inches(4.2)
        text(slide, x, Inches(2.95), Inches(3.85), Inches(0.3), label.upper(), 10, True, CYAN, PP_ALIGN.CENTER, FONT_H)
        for j, name in enumerate(names):
            rect(slide, x, Inches(3.35) + j * Inches(0.82), Inches(3.85), Inches(0.64), WHITE, line=BORDER)
            text(slide, x + Inches(0.12), Inches(3.54) + j * Inches(0.82), Inches(3.61), Inches(0.25),
                 name, 11, True, BRAND, PP_ALIGN.CENTER, FONT_H)
    rect(slide, Inches(3.55), Inches(6.15), Inches(6.25), Inches(0.58), GREEN)
    text(slide, Inches(3.7), Inches(6.32), Inches(5.95), Inches(0.25),
         "Control de Calidad Jurídica → puerta antes de la abogada", 12, True, WHITE, PP_ALIGN.CENTER, FONT_H)


def agent_detail(prs, n, agent):
    slide = blank(prs)
    header(slide, agent["name"], agent["role"], n, "Responsabilidades y skills")
    card(slide, Inches(0.7), Inches(1.82), Inches(5.75), Inches(1.3),
         "Por qué existe", agent["why"], BRAND, 12)
    card(slide, Inches(6.85), Inches(1.82), Inches(5.75), Inches(1.3),
         "Entrada → salida", agent["io"], CYAN, 12)

    rect(slide, Inches(0.7), Inches(3.42), Inches(5.75), Inches(2.65), WHITE, line=BORDER)
    text(slide, Inches(0.95), Inches(3.65), Inches(5.2), Inches(0.3),
         "RESPONSABILIDADES", 11, True, BRAND, font=FONT_H)
    rich_text(slide, Inches(0.95), Inches(4.05), Inches(5.1), Inches(1.75),
              ["• " + item for item in agent["does"]], 12, BODY, 3)

    rect(slide, Inches(6.85), Inches(3.42), Inches(5.75), Inches(2.65), WHITE, line=BORDER)
    text(slide, Inches(7.1), Inches(3.65), Inches(5.2), Inches(0.3),
         "SKILLS CLAVE (PROCEDIMIENTOS)", 11, True, BLUE, font=FONT_H)
    rich_text(slide, Inches(7.1), Inches(4.05), Inches(5.1), Inches(1.75),
              ["• " + item for item in agent["skills"]], 12, BODY, 3)

    rect(slide, Inches(0.7), Inches(6.3), Inches(11.9), Inches(0.48), SOFT)
    text(slide, Inches(0.9), Inches(6.42), Inches(11.5), Inches(0.25),
         "Límite: " + agent["limit"], 11, True, AMBER, PP_ALIGN.CENTER, FONT_H)


def daily_map(prs, n):
    slide = blank(prs)
    header(slide, "Cómo escoger una tarea", "Piense primero en el producto que necesita, no en el nombre técnico del agente.", n, "Tareas puntuales")
    rows = [
        ("Necesito entender qué pasó", "Cronología + matriz hecho–fuente", "Cronología y Hechos"),
        ("Necesito saber qué podría configurar", "Hipótesis + matriz tipo–hecho–prueba", "Tipicidad + Evidencia"),
        ("Necesito saber qué hacer ahora", "Etapa + oportunidad + ruta", "Ruta Procesal"),
        ("Tengo audiencia próxima", "Guion + solicitudes + checklist", "Audiencias + Calidad"),
        ("Necesito presentar una pieza", "Plan + borrador + revisión + firma", "Redacción + Calidad"),
        ("Necesito controlar el caso", "Estado + alerta + tarea + responsable", "Seguimiento"),
    ]
    text(slide, Inches(0.8), Inches(1.77), Inches(4.0), Inches(0.3), "NECESIDAD", 10, True, CYAN, font=FONT_H)
    text(slide, Inches(4.65), Inches(1.77), Inches(4.2), Inches(0.3), "PRODUCTO DE CIERRE", 10, True, CYAN, font=FONT_H)
    text(slide, Inches(9.1), Inches(1.77), Inches(3.3), Inches(0.3), "ÁREA PRINCIPAL", 10, True, CYAN, font=FONT_H)
    for i, (need, product, area) in enumerate(rows):
        y = Inches(2.1) + i * Inches(0.72)
        rect(slide, Inches(0.7), y, Inches(12), Inches(0.62), WHITE if i % 2 == 0 else SURFACE, line=BORDER)
        text(slide, Inches(0.9), y + Inches(0.17), Inches(3.55), Inches(0.25), need, 11, True, HEADING)
        text(slide, Inches(4.65), y + Inches(0.17), Inches(4.15), Inches(0.25), product, 11, False, BODY)
        text(slide, Inches(9.1), y + Inches(0.17), Inches(3.25), Inches(0.25), area, 11, True, BRAND)


def workflow(prs, n, title_value, subtitle, ask, required, chain, outputs, close):
    slide = blank(prs)
    header(slide, title_value, subtitle, n, "Tareas de inicio a cierre")
    card(slide, Inches(0.7), Inches(1.75), Inches(5.8), Inches(1.25),
         "Cómo comienza — pida así", "“" + ask + "”", CYAN, 13)
    card(slide, Inches(6.8), Inches(1.75), Inches(5.8), Inches(1.25),
         "Insumos mínimos", required, BRAND, 12)

    text(slide, Inches(0.75), Inches(3.25), Inches(11.8), Inches(0.25),
         "CADENA DE TRABAJO", 10, True, CYAN, font=FONT_H)
    count = len(chain)
    bw = min(2.25, (11.65 - 0.3 * (count - 1)) / count)
    total = count * bw + (count - 1) * 0.3
    start = (13.333 - total) / 2
    for i, (agent_name, skill_name) in enumerate(chain):
        x = Inches(start + i * (bw + 0.3))
        rect(slide, x, Inches(3.65), Inches(bw), Inches(1.35), NAVY if i in (0, count - 1) else BRAND)
        text(slide, x + Inches(0.08), Inches(3.88), Inches(bw - 0.16), Inches(0.38),
             agent_name, 11, True, WHITE, PP_ALIGN.CENTER, FONT_H)
        text(slide, x + Inches(0.08), Inches(4.38), Inches(bw - 0.16), Inches(0.38),
             skill_name, 9, False, SOFT, PP_ALIGN.CENTER)
        if i < count - 1:
            arrow(slide, x + Inches(bw + 0.03), Inches(4.2), Inches(0.22))

    card(slide, Inches(0.7), Inches(5.35), Inches(5.8), Inches(1.25),
         "Productos esperados", outputs, GREEN, 12)
    card(slide, Inches(6.8), Inches(5.35), Inches(5.8), Inches(1.25),
         "Cuándo termina", close, AMBER, 12)


def controls(prs, n):
    slide = blank(prs)
    header(slide, "Controles que conserva la abogada", "La automatización prepara; la responsabilidad profesional no se delega.", n, "Cierre")
    items = [
        ("Estrategia", "Define rumbo, pretensiones y prioridades."),
        ("Hechos y fuentes", "Valida lo confirmado y resuelve pendientes."),
        ("Citas y radicados", "Comprueba normas, autoridad, fechas y número."),
        ("Tono y víctima", "Evita revictimización y ajusta enfoque."),
        ("Salida externa", "Decide qué se comparte y con quién."),
        ("Firma y radicación", "Aprueba, firma y ordena seguimiento."),
    ]
    for i, (title_value, body) in enumerate(items):
        row, col = divmod(i, 3)
        card(slide, Inches(0.7) + col * Inches(4.15), Inches(1.85) + row * Inches(2.1),
             Inches(3.85), Inches(1.8), title_value, body, CYAN if i == 5 else BRAND, 13)
    rect(slide, Inches(1.4), Inches(6.18), Inches(10.55), Inches(0.58), NAVY)
    text(slide, Inches(1.6), Inches(6.34), Inches(10.15), Inches(0.25),
         "Regla de oro: la IA propone; la abogada revisa, ajusta y aprueba.",
         14, True, WHITE, PP_ALIGN.CENTER, FONT_H)


def close(prs):
    slide = blank(prs, NAVY)
    rect(slide, Inches(0), Inches(0), Inches(0.16), H, CYAN, radius=False)
    text(slide, Inches(0.9), Inches(1.0), Inches(6), Inches(0.3),
         "LEXIATEK · DESPACHO VIRTUAL", 12, True, CYAN, font=FONT_H)
    text(slide, Inches(0.9), Inches(1.55), Inches(11.4), Inches(0.75),
         "La forma correcta de trabajar", 32, True, WHITE, font=FONT_H)
    rules = [
        ("1", "Pida el producto y aporte contexto", "Objetivo, hechos, documentos, etapa y fecha relevante."),
        ("2", "Deje que el despacho distribuya", "Coordinador → especialistas → Calidad."),
        ("3", "Cierre de forma profesional", "Resuelva pendientes → ajuste → apruebe → firme → haga seguimiento."),
    ]
    for i, (num, title_value, body) in enumerate(rules):
        y = Inches(2.65) + i * Inches(1.15)
        rect(slide, Inches(0.9), y, Inches(11.55), Inches(0.92), BRAND)
        text(slide, Inches(1.17), y + Inches(0.25), Inches(0.45), Inches(0.3),
             num, 20, True, CYAN, PP_ALIGN.CENTER, FONT_H)
        text(slide, Inches(1.85), y + Inches(0.13), Inches(4.7), Inches(0.3),
             title_value, 15, True, WHITE, font=FONT_H)
        text(slide, Inches(6.0), y + Inches(0.18), Inches(5.9), Inches(0.45),
             body, 12, False, SOFT)
    text(slide, Inches(0.9), Inches(6.55), Inches(11), Inches(0.3),
         "Penal–víctimas · Colombia · Ley 906 · revisión humana obligatoria", 11, False, MUTED)


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    cover(prs)
    n = 2
    learning(prs, n); n += 1
    dedicated_agents(prs, n); n += 1
    agent_skill_tool(prs, n); n += 1
    skill_anatomy(prs, n); n += 1
    tools_slide(prs, n); n += 1
    despacho(prs, n); n += 1
    architecture(prs, n); n += 1
    task_lifecycle(prs, n); n += 1
    agents_map(prs, n); n += 1
    for agent in AGENTS:
        agent_detail(prs, n, agent)
        n += 1
    daily_map(prs, n); n += 1
    workflow(
        prs, n,
        "Tarea 1 · Abrir un caso y ordenar hechos",
        "El objetivo es producir una base factual utilizable, no una conclusión penal apresurada.",
        "Tengo una nueva víctima. Ordene los hechos, identifique urgencia y diga qué falta.",
        "Relato inicial · fecha del último hecho · soportes · autoridad/radicado si existe",
        [
            ("Coordinador", "urgencia + faltantes"),
            ("Cronología", "línea de tiempo"),
            ("Víctimas", "intereses + riesgos"),
            ("Calidad", "hechos soportados"),
        ],
        "Cronología con fuentes · vacíos y preguntas · urgencia · objetivo preliminar",
        "Cuando la abogada valida la cronología y decide el siguiente análisis o actuación.",
    ); n += 1
    workflow(
        prs, n,
        "Tarea 2 · Analizar tipicidad preliminar",
        "La salida es una hipótesis de trabajo con soporte y riesgos, nunca una conclusión definitiva.",
        "Mapee estos hechos contra posibles tipos penales y marque lo que falta probar.",
        "Cronología validada · actores · conductas · soportes disponibles",
        [
            ("Tipicidad", "tipo–hecho–prueba"),
            ("Evidencia", "brechas probatorias"),
            ("Calidad", "soporte + atipicidad"),
        ],
        "Hipótesis preliminares · matriz elemento–hecho–prueba · riesgos · preguntas",
        "Cuando la abogada acepta, ajusta o descarta hipótesis para definir estrategia.",
    ); n += 1
    workflow(
        prs, n,
        "Tarea 3 · Preparar una audiencia",
        "La preparación une etapa procesal, objetivo de víctima, prueba y táctica oral.",
        "Prepare guion, solicitudes, riesgos y checklist para la audiencia de mañana.",
        "Tipo y fecha de audiencia · objetivo · última actuación · prueba · riesgos de víctima",
        [
            ("Ruta 906", "etapa + oportunidad"),
            ("Audiencias", "guion + solicitudes"),
            ("Víctimas", "no revictimizar"),
            ("Calidad", "coherencia + riesgos"),
        ],
        "Resumen ejecutivo · guion · solicitudes · contraargumentos · checklist",
        "Cuando la abogada ensaya, ajusta el guion y llega con decisiones y documentos listos.",
    ); n += 1
    workflow(
        prs, n,
        "Tarea 4 · Redactar un memorial de impulso",
        "Una pieza escrita exige plan aprobado, insumos verificados y firma humana.",
        "Prepare memorial de impulso por inactividad, con petición concreta.",
        "Radicado · autoridad · última actuación · fecha · objetivo · anexos soporte",
        [
            ("Seguimiento", "inactividad + estado"),
            ("Ruta 906", "oportunidad"),
            ("Redacción", "hechos + fundamentos"),
            ("Calidad", "citas + soporte"),
        ],
        "Borrador estructurado · fuentes · anexos · pendientes claramente marcados",
        "Cuando la abogada resuelve pendientes, ajusta, aprueba, firma y ordena radicación.",
    ); n += 1
    workflow(
        prs, n,
        "Tarea 5 · Cerrar con seguimiento",
        "Toda actuación debe dejar estado, responsable, fecha y próximo control.",
        "Registre lo radicado, cree alertas y prepare el próximo control del expediente.",
        "Pieza final · constancia de radicación · fecha · autoridad · término esperado",
        [
            ("Seguimiento", "registro + alertas"),
            ("Coordinador", "tarea + responsable"),
            ("Calidad", "resumen operativo"),
        ],
        "Actuación registrada · alerta · tarea asignada · reporte interno y próximo control",
        "Cuando existe trazabilidad y una fecha concreta de revisión; no al enviar el documento.",
    ); n += 1
    controls(prs, n)
    close(prs)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(f"Wrote {OUT} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    build()
